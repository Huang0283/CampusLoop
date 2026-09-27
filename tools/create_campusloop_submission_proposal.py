from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    HRFlowable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
DOCX_OUT = ROOT / "CampusLoop_Project_Proposal_Submission.docx"
PDF_OUT = ROOT / "output" / "pdf" / "CampusLoop_Project_Proposal_Submission.pdf"

GROUP_NUMBER = "5"
PROGRAMME = "BSc (AI)"
PROPOSAL_DATE = "27 September 2026"
ESTIMATED_PERIOD = "28 September 2026 - 1 November 2026"

MEMBERS = [
    ("Huang Qiwei", "黄麒玮", "20253801012"),
    ("Chen Zihong", "陈梓弘", "20253801098"),
    ("Gao Junjie", "高峻杰", "20253801010"),
    ("Hu Keming", "胡可铭", "20253801017"),
    ("Li Mingyuan", "李明远", "20253801016"),
    ("Ma Xuan", "马萱", "20253801040"),
    ("Chen Hongqing", "陈红卿", "20253801080"),
    ("Luo Kangrui", "骆康睿", "20253801022"),
    ("Liang Jiawen", "梁迦文", "20253801026"),
    ("Ye Zidan", "叶子聃", "20253801096"),
]

TIMELINE_ROWS = [
    (
        "Phase 1 / Week 1",
        "28 Sep - 4 Oct",
        "Confirm scope, roles, journeys, domain and AI feasibility, risks and tests.",
        "Approved Phase 1 baseline.",
    ),
    (
        "Phase 2 / Week 2",
        "5 - 11 Oct",
        "Create clickable prototypes; freeze API, SDK, database, environment, AI and test contracts.",
        "Runnable contract baseline.",
    ),
    (
        "Phase 3 / Week 3",
        "12 - 18 Oct",
        "Implement authentication, marketplace and the real two-account transaction flow.",
        "Persisted MVP loop.",
    ),
    (
        "Phase 4 / Week 4",
        "19 - 25 Oct",
        "Integrate two justified AI functions; test metrics, explanation, privacy and fallback; freeze.",
        "Frozen release candidate.",
    ),
    (
        "Phase 5 / Week 5",
        "26 Oct - 1 Nov",
        "Run regression, security, load, recovery and clean deployment; finish documents.",
        "Tagged submission package.",
    ),
]

ACTION_ROWS = [
    (
        "M1 - Huang Qiwei",
        "Requirements, scope, traceability, integration gates, risks, final report and presentation.",
        "4 Oct; weekly; 1 Nov",
        "Planned",
    ),
    (
        "M2 - Chen Zihong",
        "Frontend shell, navigation, auth/profile, shared components and page states.",
        "11 Oct; 18 Oct",
        "Planned",
    ),
    (
        "M3 - Gao Junjie",
        "Marketplace, listings, wanted requests, search, matching and price-advice UI.",
        "11 Oct; 18 Oct; 25 Oct",
        "Planned",
    ),
    (
        "M4 - Hu Keming",
        "Chat, offers, orders, meetup, completion, reviews, reports and notifications UI.",
        "11 Oct; 18 Oct",
        "Planned",
    ),
    (
        "M5 - Li Mingyuan",
        "Authentication, sessions, privacy, roles, field visibility and administrator permissions.",
        "11 Oct; 18 Oct",
        "Planned",
    ),
    (
        "M6 - Ma Xuan",
        "Marketplace/transaction entities, APIs, WebSocket, states, transactions and idempotency.",
        "11 Oct; 18 Oct; 25 Oct",
        "Planned",
    ),
    (
        "M7 - Chen Hongqing",
        "Keyword/semantic retrieval, wanted matching, explanations and search evaluation.",
        "11 Oct; 18 Oct; 25 Oct",
        "Planned",
    ),
    (
        "M8 - Luo Kangrui",
        "Price reference, trust evidence, risk-review assistance, evaluation and fallback.",
        "11 Oct; 18 Oct; 25 Oct",
        "Planned",
    ),
    (
        "M9 - Liang Jiawen",
        "Database, migrations, seed data, environment, storage, jobs, CI, deployment and recovery.",
        "11 Oct; 18 Oct; 1 Nov",
        "Planned",
    ),
    (
        "M10 - Ye Zidan",
        "Independent QA, defects, permission/E2E tests, load/recovery tests and release sign-off.",
        "Weekly; 1 Nov",
        "Planned",
    ),
]

OBJECTIVES = [
    ("O1", "Define and implement account, profile, role, session and privacy boundaries with explicit authorised and unauthorised results."),
    ("O2", "Provide product and wanted-request publishing, browsing, keyword search, filtering, favourites and lifecycle management through a responsive interface."),
    ("O3", "Demonstrate a real two-account transaction from chat and structured negotiation through order creation, meetup agreement, independent completion confirmation and review."),
    ("O4", "Integrate at least two justified intelligent functions with a simple baseline, versioned input and output, explanation, evaluation measure and unavailable-service fallback."),
    ("O5", "Provide contextual reporting and administrator review while preventing automated risk signals from directly changing accounts, orders or penalties."),
    ("O6", "Deliver a reproducible engineering package containing OpenAPI, generated SDK, migrations, seed data, tests, CI evidence, clean-deployment instructions, backup and recovery guidance, technical documentation and a final demonstration."),
]

BENEFIT_ROWS = [
    ("Student buyers", "Faster discovery, persistent wanted requests, clearer price references, structured negotiation and safer coordination."),
    ("Student sellers", "Standardised listings, access to relevant buyers, controlled listing states, traceable offers and evidence from completed transactions."),
    ("Campus community", "Lower reuse costs, reduced waste, clearer behaviour expectations and a practical local circular-economy channel."),
    ("Administrators", "Contextual reports, controlled evidence access, auditable review decisions and clear separation between allegations, automated signals and human outcomes."),
    ("Project team", "Experience across requirements, architecture, frontend, backend, databases, AI evaluation, testing, deployment, documentation and team integration."),
]


def register_pdf_fonts():
    font_dir = Path(r"C:\Windows\Fonts")
    pdfmetrics.registerFont(TTFont("TimesNewRoman", str(font_dir / "times.ttf")))
    pdfmetrics.registerFont(TTFont("TimesNewRoman-Bold", str(font_dir / "timesbd.ttf")))
    pdfmetrics.registerFont(TTFont("TimesNewRoman-Italic", str(font_dir / "timesi.ttf")))
    pdfmetrics.registerFont(TTFont("TimesNewRoman-BoldItalic", str(font_dir / "timesbi.ttf")))
    pdfmetrics.registerFont(TTFont("SimSun", str(font_dir / "simsun.ttc"), subfontIndex=0))


def pdf_styles():
    styles = getSampleStyleSheet()
    common = dict(fontName="TimesNewRoman", fontSize=12, leading=18, textColor=colors.black)
    return {
        "body": ParagraphStyle(
            "Body",
            parent=styles["Normal"],
            alignment=TA_JUSTIFY,
            spaceAfter=8,
            **common,
        ),
        "body_left": ParagraphStyle(
            "BodyLeft",
            parent=styles["Normal"],
            alignment=TA_LEFT,
            spaceAfter=8,
            **common,
        ),
        "chapter": ParagraphStyle(
            "Chapter",
            parent=styles["Heading1"],
            fontName="TimesNewRoman-Bold",
            fontSize=12,
            leading=18,
            alignment=TA_LEFT,
            spaceAfter=14,
            keepWithNext=True,
        ),
        "subheading": ParagraphStyle(
            "Subheading",
            parent=styles["Heading2"],
            fontName="TimesNewRoman-Bold",
            fontSize=12,
            leading=18,
            alignment=TA_LEFT,
            spaceBefore=8,
            spaceAfter=6,
            keepWithNext=True,
        ),
        "cover_center": ParagraphStyle(
            "CoverCenter",
            parent=styles["Normal"],
            fontName="TimesNewRoman",
            fontSize=12,
            leading=18,
            alignment=TA_CENTER,
            spaceAfter=8,
        ),
        "cover_bold": ParagraphStyle(
            "CoverBold",
            parent=styles["Normal"],
            fontName="TimesNewRoman-Bold",
            fontSize=12,
            leading=18,
            alignment=TA_CENTER,
            spaceAfter=8,
        ),
        "table": ParagraphStyle(
            "Table",
            parent=styles["Normal"],
            fontName="TimesNewRoman",
            fontSize=12,
            leading=18,
            alignment=TA_LEFT,
        ),
        "table_bold": ParagraphStyle(
            "TableBold",
            parent=styles["Normal"],
            fontName="TimesNewRoman-Bold",
            fontSize=12,
            leading=18,
            alignment=TA_LEFT,
        ),
        "name": ParagraphStyle(
            "Name",
            parent=styles["Normal"],
            fontName="TimesNewRoman",
            fontSize=12,
            leading=18,
            alignment=TA_LEFT,
        ),
    }


def p(text, style):
    return Paragraph(escape(text), style)


def rich(text, style):
    return Paragraph(text, style)


def chapter_title(number, title, styles):
    return [Paragraph(f"{number}  {escape(title)}", styles["chapter"]), HRFlowable(width="100%", thickness=0.8, color=colors.black, spaceAfter=12)]


def bullet(text, styles):
    return Paragraph(f"&#8226;&nbsp;&nbsp;{escape(text)}", styles["body_left"])


def body_footer(canvas, doc):
    if canvas.getPageNumber() == 1:
        return
    canvas.saveState()
    canvas.setFont("TimesNewRoman", 12)
    canvas.drawCentredString(A4[0] / 2, 0.55 * inch, str(canvas.getPageNumber() - 1))
    canvas.restoreState()


def pdf_table(data, widths, styles, header=True, repeat_rows=1):
    rows = []
    for r_idx, row in enumerate(data):
        cells = []
        for value in row:
            if isinstance(value, Paragraph):
                cells.append(value)
            else:
                cells.append(p(str(value), styles["table_bold"] if header and r_idx == 0 else styles["table"]))
        rows.append(cells)
    table = Table(rows, colWidths=widths, repeatRows=repeat_rows if header else 0, hAlign="LEFT")
    commands = [
        ("FONTNAME", (0, 0), (-1, -1), "TimesNewRoman"),
        ("FONTSIZE", (0, 0), (-1, -1), 12),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("LINEABOVE", (0, 0), (-1, 0), 0.8, colors.black),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.black),
        ("LINEBELOW", (0, -1), (-1, -1), 0.8, colors.black),
        ("BOX", (0, 0), (-1, -1), 0.0, colors.white),
    ]
    for row_index in range(1, len(rows) - 1):
        commands.append(("LINEBELOW", (0, row_index), (-1, row_index), 0.25, colors.HexColor("#777777")))
    table.setStyle(TableStyle(commands))
    return table


def build_pdf():
    register_pdf_fonts()
    styles = pdf_styles()
    PDF_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(PDF_OUT),
        pagesize=A4,
        leftMargin=inch,
        rightMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
        title="CampusLoop AI Project Proposal",
        author="CS Group 5",
        subject="Introduction to Software Engineering Assessment",
    )
    story = []

    # Cover page - restrained OceanScope-inspired hierarchy and spacing.
    story += [Spacer(1, 28), p("CAMPUSLOOP AI", styles["cover_bold"]), Spacer(1, 8)]
    story += [p("An AI-Enhanced Trusted Campus Second-Hand Marketplace", styles["cover_center"]), Spacer(1, 24)]
    story += [p("PROJECT PROPOSAL", styles["cover_bold"]), p("Introduction to Software Engineering Assessment", styles["cover_center"]), Spacer(1, 16)]
    metadata = [
        ["Group Number", GROUP_NUMBER],
        ["Programme of Study", PROGRAMME],
        ["Team Size", "10 students"],
        ["Proposal Date", PROPOSAL_DATE],
        ["Estimated Project Period", ESTIMATED_PERIOD],
    ]
    story.append(pdf_table(metadata, [1.9 * inch, 3.9 * inch], styles, header=False, repeat_rows=0))
    story.append(Spacer(1, 10))
    member_data = [["Group Member", "Student ID", "Group Member", "Student ID"]]
    for index in range(5):
        left = MEMBERS[index]
        right = MEMBERS[index + 5]
        left_name = rich(f"{escape(left[0])} / <font name='SimSun'>{escape(left[1])}</font>", styles["name"])
        right_name = rich(f"{escape(right[0])} / <font name='SimSun'>{escape(right[1])}</font>", styles["name"])
        member_data.append([left_name, left[2], right_name, right[2]])
    story.append(pdf_table(member_data, [1.75 * inch, 1.15 * inch, 1.75 * inch, 1.15 * inch], styles))
    story.append(Spacer(1, 10))
    story.append(p("This document describes planned work. It does not claim that any proposed feature, integration or evaluation result has already been completed.", styles["cover_center"]))
    story.append(PageBreak())

    # Page 1 of body.
    story += chapter_title("1", "Project Title and Problem", styles)
    story.append(p("Project title: CampusLoop AI - An AI-Enhanced Trusted Campus Second-Hand Marketplace", styles["subheading"]))
    story.append(p("CampusLoop AI is a proposed web platform for trusted second-hand trading within a university community. It combines a conventional marketplace and a traceable transaction process with explainable assistance for search, matching and price reference. The basic marketplace remains usable when optional AI services are unavailable.", styles["body"]))
    story.append(p("Problem", styles["subheading"]))
    story.append(p("Students currently buy and sell textbooks, electronics, furniture and daily necessities through fragmented channels such as messaging groups, social feeds and general marketplaces. Listings use inconsistent descriptions, become outdated quickly and are difficult to search. Buyers may describe a purpose and budget while sellers describe a brand or model, so keyword-only discovery can miss relevant items.", styles["body"]))
    story.append(p("After discovering an item, students still need to negotiate, agree a meeting time and place, confirm the handover and resolve unsuccessful interactions. Informal channels provide limited support for structured offers, transaction status, two-party confirmation, reviews, contextual reports or auditable administrator decisions. This creates unnecessary coordination cost and weakens trust.", styles["body"]))
    story.append(p("The project is important because a campus marketplace can reduce student living costs, extend the useful life of goods and support a local circular economy. However, these benefits require clear permissions, privacy protection, transaction consistency and honest AI limitations rather than a visually impressive but unreliable demonstration.", styles["body"]))
    story.append(p("The proposed course scope excludes real payment, delivery, university identity-provider integration and automatic punitive decisions. Campus identity is simulated for demonstration, and any risk signal is routed to human review.", styles["body"]))
    story.append(PageBreak())

    # Page 2 of body.
    story += chapter_title("2", "Proposed Solution", styles)
    story.append(p("CampusLoop AI will provide one browser-based workflow for public discovery, authenticated trading and administrator governance. Guests may browse public marketplace and wanted information. Registration or login is required only for actions that depend on identity, including publishing, editing, favouriting, chatting, negotiating, ordering, reviewing and reporting.", styles["body"]))
    story.append(p("Core solution components", styles["subheading"]))
    for item in [
        "Account and permission layer: registration, login, profile, student and administrator roles, session expiry and field-level privacy.",
        "Marketplace and wanted requests: listings, images, categories, condition, price, search, filters, favourites, seller management, expiry and lifecycle states.",
        "Transaction coordination: contextual chat, structured offers and counter-offers, order creation, meetup versions, two-party completion, reviews and notifications.",
        "Governance: contextual reporting, evidence visibility, administrator review, reasoned decisions, appeals and audit fields.",
        "Explainable assistance: keyword and semantic retrieval, wanted-to-product matching, price-reference ranges, transparent factors and deterministic non-AI fallback.",
        "Engineering foundation: React and TypeScript frontend; FastAPI backend; PostgreSQL and optional pgvector; generated OpenAPI SDK; Redis, object storage and workers where justified; CI, logs, backup and recovery instructions.",
    ]:
        story.append(bullet(item, styles))
    story.append(p("The backend and database will remain the sources of truth for identity, permission, listing availability, accepted offer, order state, meetup version, completion confirmation and review eligibility. Frontend state and AI suggestions are derived views and cannot override authoritative business facts.", styles["body"]))
    story.append(p("The system will be developed as a modular proof of concept. Optional services must have visible unavailable states and safe fallbacks so that browsing and the basic transaction loop do not fail solely because an intelligent function, cache or worker is unavailable.", styles["body"]))
    story.append(PageBreak())

    # Page 3 of body.
    story += chapter_title("3", "Objectives", styles)
    obj_data = [["ID", "Specific and achievable objective"]] + list(OBJECTIVES)
    story.append(pdf_table(obj_data, [0.55 * inch, 5.25 * inch], styles))
    story.append(Spacer(1, 14))
    story.append(p("Acceptance basis", styles["subheading"]))
    story.append(p("Objectives will be accepted only from an integrated repository commit. A feature author performs initial checks, another group member reviews the pull request, M10 independently executes the relevant acceptance scenarios and M1 confirms that the result remains within the agreed scope. Mock data may support Week 2 prototypes but cannot close a Week 3 real-integration objective.", styles["body"]))
    story.append(p("The primary demonstration will use a buyer and seller account to publish, discover, discuss, negotiate, confirm a meetup, complete the exchange, review one another and create a contextual report. Permission, invalid-state, duplicate-action and service-unavailable scenarios will also be shown.", styles["body"]))
    story.append(PageBreak())

    # Page 4 of body.
    story += chapter_title("4", "Expected Benefits", styles)
    benefit_data = [["Stakeholder", "Expected benefit"]] + [list(row) for row in BENEFIT_ROWS]
    story.append(pdf_table(benefit_data, [1.45 * inch, 4.35 * inch], styles))
    story.append(Spacer(1, 14))
    story.append(p("Expected indicators of success", styles["subheading"]))
    for item in [
        "A non-author can start the documented environment and apply migrations without private instructions from the original developer.",
        "Two accounts can complete the agreed MVP transaction using real APIs and persistent data.",
        "Non-participants cannot access private conversations, orders or report evidence.",
        "Accepted AI functions have reproducible baseline comparisons, explanations and safe fallback behaviour.",
        "The final release has no unresolved P0 or P1 defect and includes regression, load, recovery and known-limitation evidence.",
    ]:
        story.append(bullet(item, styles))
    story.append(p("The project will not claim commercial readiness, guaranteed fraud detection, guaranteed price accuracy or unrestricted production capacity. Benefits will be reported only to the extent supported by executed evidence.", styles["body"]))
    story.append(PageBreak())

    # Page 5 of body.
    story += chapter_title("5", "Timeline", styles)
    story.append(p("The project uses five consecutive one-week phases from 28 September 2026. The dates are estimates and the official MyAberdeen deadline takes precedence. Each gate requires a reviewable increment, evidence and handoff; incomplete work keeps an owner and recovery action.", styles["body"]))
    timeline_data = [["Phase", "Estimated dates", "Main activities", "Expected completion"]] + [list(row) for row in TIMELINE_ROWS]
    story.append(pdf_table(timeline_data, [1.05 * inch, 1.0 * inch, 2.35 * inch, 1.4 * inch], styles))
    story.append(PageBreak())

    # Pages 6 and 7 of body.
    story += chapter_title("6", "Action Plan", styles)
    story.append(p("The action plan assigns a primary owner to each work area. Progress is shown as Planned because this proposal describes the agreed work programme and does not claim that a phase has already been completed.", styles["body"]))
    action_header = [["Responsible member", "Required activity", "Deadline", "Status"]]
    story.append(pdf_table(action_header + [list(row) for row in ACTION_ROWS[:5]], [1.2 * inch, 2.55 * inch, 1.15 * inch, 0.9 * inch], styles))
    story.append(Spacer(1, 10))
    story.append(p("Cross-group dependency order", styles["subheading"]))
    story.append(p("Rules and contracts are published before consumers integrate them: M5/M6 provide identity and business contracts; M7/M8 provide versioned intelligent-function outputs; M9 provides the environment and migrations; M2-M4 integrate the user experience; M10 independently tests the integrated result; M1 controls scope and the phase handoff.", styles["body"]))
    story.append(PageBreak())

    story += chapter_title("6", "Action Plan - continued", styles)
    story.append(pdf_table(action_header + [list(row) for row in ACTION_ROWS[5:]], [1.2 * inch, 2.55 * inch, 1.15 * inch, 0.9 * inch], styles))
    story.append(Spacer(1, 8))
    story.append(p("Academic integrity and AI use", styles["subheading"]))
    story.append(p("This proposal represents Group 5's planned work. Generative AI assisted with structure, wording, formatting and consistency checks. The group will review and approve the final submission and remains responsible for its accuracy and originality. Any external text, code, figure, image, data or idea used during the project will be acknowledged according to University guidance. This proposal contains no external figure or unverified result.", styles["body"]))
    story.append(p("Supervisor approval", styles["subheading"]))
    story.append(p("The group will discuss the scope, five-week schedule and technical approach with the assigned supervisor before submission. Any substantial later change will also be discussed and recorded.", styles["body"]))
    doc.build(story, onFirstPage=body_footer, onLaterPages=body_footer)


def set_run_font(run, bold=None, italic=None, east_asia="Times New Roman"):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(12)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_doc_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, before=0):
    paragraph.alignment = alignment
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)


def add_doc_paragraph(doc, text, bold=False, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6, before=0):
    paragraph = doc.add_paragraph()
    set_doc_paragraph(paragraph, alignment=alignment, after=after, before=before)
    run = paragraph.add_run(text)
    set_run_font(run, bold=bold)
    return paragraph


def add_doc_chapter(doc, number, title):
    paragraph = add_doc_paragraph(doc, f"{number}  {title}", bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT, after=6)
    paragraph.paragraph_format.keep_with_next = True
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "000000")
    borders.append(bottom)
    p_pr.append(borders)


def add_doc_subheading(doc, text):
    paragraph = add_doc_paragraph(doc, text, bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT, after=4, before=6)
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_doc_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    set_doc_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.LEFT, after=3)
    for run in paragraph.runs:
        set_run_font(run)
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def prevent_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:cantSplit")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def set_table_edges(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, size in (("top", "8"), ("bottom", "8"), ("insideH", "4")):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), size)
        item.set(qn("w:space"), "0")
        item.set(qn("w:color"), "000000" if edge != "insideH" else "777777")
        borders.append(item)
    for edge in ("left", "right", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "nil")
        borders.append(item)
    tbl_pr.append(borders)


def add_doc_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_edges(table)
    header = table.rows[0]
    set_repeat_header(header)
    prevent_split(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.width = widths[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        set_doc_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.LEFT, after=0)
        run = paragraph.add_run(value)
        set_run_font(run, bold=True)
    for values in rows:
        row = table.add_row()
        prevent_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = widths[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            paragraph = cell.paragraphs[0]
            set_doc_paragraph(paragraph, alignment=WD_ALIGN_PARAGRAPH.LEFT, after=0)
            run = paragraph.add_run(str(value))
            set_run_font(run, east_asia="SimSun" if any("\u4e00" <= ch <= "\u9fff" for ch in str(value)) else "Times New Roman")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_number(section):
    section.footer.is_linked_to_previous = False
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])
    set_run_font(run)
    sect_pr = section._sectPr
    page_num = OxmlElement("w:pgNumType")
    page_num.set(qn("w:start"), "1")
    sect_pr.append(page_num)


def configure_section(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.45)


def build_docx():
    doc = Document()
    configure_section(doc.sections[0])
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing = 1.5

    # Cover page.
    add_doc_paragraph(doc, "", alignment=WD_ALIGN_PARAGRAPH.CENTER, after=24)
    add_doc_paragraph(doc, "CAMPUSLOOP AI", bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    add_doc_paragraph(doc, "An AI-Enhanced Trusted Campus Second-Hand Marketplace", alignment=WD_ALIGN_PARAGRAPH.CENTER, after=30)
    add_doc_paragraph(doc, "PROJECT PROPOSAL", bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    add_doc_paragraph(doc, "Introduction to Software Engineering Assessment", alignment=WD_ALIGN_PARAGRAPH.CENTER, after=20)
    add_doc_table(
        doc,
        ["Item", "Details"],
        [
            ["Group Number", GROUP_NUMBER],
            ["Programme of Study", PROGRAMME],
            ["Team Size", "10 students"],
            ["Proposal Date", PROPOSAL_DATE],
            ["Estimated Project Period", ESTIMATED_PERIOD],
        ],
        [Inches(1.8), Inches(4.3)],
    )
    member_rows = []
    for index in range(5):
        left = MEMBERS[index]
        right = MEMBERS[index + 5]
        member_rows.append([f"{left[0]} / {left[1]}", left[2], f"{right[0]} / {right[1]}", right[2]])
    add_doc_table(doc, ["Group Member", "Student ID", "Group Member", "Student ID"], member_rows, [Inches(1.9), Inches(1.1), Inches(1.9), Inches(1.1)])
    add_doc_paragraph(
        doc,
        "This document describes planned work. It does not claim that any proposed feature, integration or evaluation result has already been completed.",
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        before=6,
    )

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(body_section)
    add_page_number(body_section)

    # Body page 1.
    add_doc_chapter(doc, "1", "Project Title and Problem")
    add_doc_subheading(doc, "Project title: CampusLoop AI - An AI-Enhanced Trusted Campus Second-Hand Marketplace")
    add_doc_paragraph(doc, "CampusLoop AI is a proposed web platform for trusted second-hand trading within a university community. It combines a conventional marketplace and a traceable transaction process with explainable assistance for search, matching and price reference. The basic marketplace remains usable when optional AI services are unavailable.")
    add_doc_subheading(doc, "Problem")
    add_doc_paragraph(doc, "Students currently buy and sell textbooks, electronics, furniture and daily necessities through fragmented channels such as messaging groups, social feeds and general marketplaces. Listings use inconsistent descriptions, become outdated quickly and are difficult to search. Buyers may describe a purpose and budget while sellers describe a brand or model, so keyword-only discovery can miss relevant items.")
    add_doc_paragraph(doc, "After discovering an item, students still need to negotiate, agree a meeting time and place, confirm the handover and resolve unsuccessful interactions. Informal channels provide limited support for structured offers, transaction status, two-party confirmation, reviews, contextual reports or auditable administrator decisions. This creates unnecessary coordination cost and weakens trust.")
    add_doc_paragraph(doc, "The project is important because a campus marketplace can reduce student living costs, extend the useful life of goods and support a local circular economy. However, these benefits require clear permissions, privacy protection, transaction consistency and honest AI limitations rather than a visually impressive but unreliable demonstration.")
    add_doc_paragraph(doc, "The proposed course scope excludes real payment, delivery, university identity-provider integration and automatic punitive decisions. Campus identity is simulated for demonstration, and any risk signal is routed to human review.")
    # Body page 2.
    add_doc_chapter(doc, "2", "Proposed Solution")
    add_doc_paragraph(doc, "CampusLoop AI will provide one browser-based workflow for public discovery, authenticated trading and administrator governance. Guests may browse public marketplace and wanted information. Registration or login is required only for actions that depend on identity, including publishing, editing, favouriting, chatting, negotiating, ordering, reviewing and reporting.")
    add_doc_subheading(doc, "Core solution components")
    for item in [
        "Account and permission layer: registration, login, profile, student and administrator roles, session expiry and field-level privacy.",
        "Marketplace and wanted requests: listings, images, categories, condition, price, search, filters, favourites, seller management, expiry and lifecycle states.",
        "Transaction coordination: contextual chat, structured offers and counter-offers, order creation, meetup versions, two-party completion, reviews and notifications.",
        "Governance: contextual reporting, evidence visibility, administrator review, reasoned decisions, appeals and audit fields.",
        "Explainable assistance: keyword and semantic retrieval, wanted-to-product matching, price-reference ranges, transparent factors and deterministic non-AI fallback.",
        "Engineering foundation: React and TypeScript frontend; FastAPI backend; PostgreSQL and optional pgvector; generated OpenAPI SDK; Redis, object storage and workers where justified; CI, logs, backup and recovery instructions.",
    ]:
        add_doc_bullet(doc, item)
    add_doc_paragraph(doc, "The backend and database will remain the sources of truth for identity, permission, listing availability, accepted offer, order state, meetup version, completion confirmation and review eligibility. Frontend state and AI suggestions are derived views and cannot override authoritative business facts.")
    add_doc_paragraph(doc, "The system will be developed as a modular proof of concept. Optional services must have visible unavailable states and safe fallbacks so that browsing and the basic transaction loop do not fail solely because an intelligent function, cache or worker is unavailable.")
    # Body page 3.
    add_doc_chapter(doc, "3", "Objectives")
    add_doc_table(doc, ["ID", "Specific and achievable objective"], OBJECTIVES, [Inches(0.6), Inches(5.5)])
    add_doc_subheading(doc, "Acceptance basis")
    add_doc_paragraph(doc, "Objectives will be accepted only from an integrated repository commit. A feature author performs initial checks, another group member reviews the pull request, M10 independently executes the relevant acceptance scenarios and M1 confirms that the result remains within the agreed scope. Mock data may support Week 2 prototypes but cannot close a Week 3 real-integration objective.")
    add_doc_paragraph(doc, "The primary demonstration will use a buyer and seller account to publish, discover, discuss, negotiate, confirm a meetup, complete the exchange, review one another and create a contextual report. Permission, invalid-state, duplicate-action and service-unavailable scenarios will also be shown.")
    # Body page 4.
    add_doc_chapter(doc, "4", "Expected Benefits")
    add_doc_table(doc, ["Stakeholder", "Expected benefit"], BENEFIT_ROWS, [Inches(1.5), Inches(4.6)])
    add_doc_subheading(doc, "Expected indicators of success")
    for item in [
        "A non-author can start the documented environment and apply migrations without private instructions from the original developer.",
        "Two accounts can complete the agreed MVP transaction using real APIs and persistent data.",
        "Non-participants cannot access private conversations, orders or report evidence.",
        "Accepted AI functions have reproducible baseline comparisons, explanations and safe fallback behaviour.",
        "The final release has no unresolved P0 or P1 defect and includes regression, load, recovery and known-limitation evidence.",
    ]:
        add_doc_bullet(doc, item)
    add_doc_paragraph(doc, "The project will not claim commercial readiness, guaranteed fraud detection, guaranteed price accuracy or unrestricted production capacity. Benefits will be reported only to the extent supported by executed evidence.")
    # Body page 5.
    add_doc_chapter(doc, "5", "Timeline")
    add_doc_paragraph(doc, "The project uses five consecutive one-week phases from 28 September 2026. The dates are estimates and the official MyAberdeen deadline takes precedence. Each gate requires a reviewable increment, evidence and handoff; incomplete work keeps an owner and recovery action.")
    add_doc_table(doc, ["Phase", "Estimated dates", "Main activities", "Expected completion"], TIMELINE_ROWS, [Inches(1.0), Inches(1.0), Inches(2.6), Inches(1.5)])
    # Body pages 6 and 7.
    add_doc_chapter(doc, "6", "Action Plan")
    add_doc_paragraph(doc, "The action plan assigns a primary owner to each work area. Progress is shown as Planned because this proposal describes the agreed work programme and does not claim that a phase has already been completed.")
    add_doc_table(doc, ["Responsible member", "Required activity", "Deadline", "Status"], ACTION_ROWS[:5], [Inches(1.25), Inches(2.65), Inches(1.2), Inches(0.9)])
    add_doc_subheading(doc, "Cross-group dependency order")
    add_doc_paragraph(doc, "Rules and contracts are published before consumers integrate them: M5/M6 provide identity and business contracts; M7/M8 provide versioned intelligent-function outputs; M9 provides the environment and migrations; M2-M4 integrate the user experience; M10 independently tests the integrated result; M1 controls scope and the phase handoff.")
    add_doc_chapter(doc, "6", "Action Plan - continued")
    add_doc_table(doc, ["Responsible member", "Required activity", "Deadline", "Status"], ACTION_ROWS[5:], [Inches(1.25), Inches(2.65), Inches(1.2), Inches(0.9)])
    add_doc_subheading(doc, "Academic integrity and AI use")
    add_doc_paragraph(doc, "This proposal represents Group 5's planned work. Generative AI assisted with structure, wording, formatting and consistency checks. The group will review and approve the final submission and remains responsible for its accuracy and originality. Any external text, code, figure, image, data or idea used during the project will be acknowledged according to University guidance. This proposal contains no external figure or unverified result.")
    add_doc_subheading(doc, "Supervisor approval")
    add_doc_paragraph(doc, "The group will discuss the scope, five-week schedule and technical approach with the assigned supervisor before submission. Any substantial later change will also be discussed and recorded.")
    doc.core_properties.title = "CampusLoop AI Project Proposal"
    doc.core_properties.subject = "Introduction to Software Engineering Assessment"
    doc.core_properties.author = "CS Group 5"
    doc.core_properties.keywords = "CampusLoop AI, software engineering, project proposal"
    doc.save(DOCX_OUT)


if __name__ == "__main__":
    build_pdf()
    build_docx()
    print(DOCX_OUT)
    print(PDF_OUT)
