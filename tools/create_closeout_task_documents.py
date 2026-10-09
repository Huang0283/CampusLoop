"""Build the two Chinese task documents from reviewable Markdown sources.

Requires python-docx. Optional --wechat-directory updates the supplied original
only after saving a recoverable backup, and also writes the eight-person plan.
"""

from __future__ import annotations

import argparse
from pathlib import Path
import re
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCES = (
    ("closeout-task-packages.md", "CampusLoop_Closeout_Task_Packages.docx"),
    (
        "eight-member-self-selected-deliverables.md",
        "CampusLoop_Eight_Member_Deliverables.docx",
    ),
)


def apply_font(style, size: int, bold: bool = False) -> None:
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")


def add_runs(paragraph, text: str) -> None:
    for index, part in enumerate(re.split(r"(`[^`]+`)", text)):
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        else:
            paragraph.add_run(part)


def build(source: Path, output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.54)
    section.left_margin = section.right_margin = Cm(2.54)
    normal = document.styles["Normal"]
    apply_font(normal, 12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [("Title", 20), ("Heading 1", 16), ("Heading 2", 13)]:
        style = document.styles[name]
        apply_font(style, size, True)
        style.font.color.rgb = RGBColor.from_string("20374A")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
    apply_font(document.styles["List Bullet"], 12)

    for line in source.read_text(encoding="utf-8").splitlines():
        if not line:
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(line[2:], "Title")
        elif line.startswith("## "):
            paragraph = document.add_paragraph(line[3:], "Heading 1")
        elif line.startswith("### "):
            paragraph = document.add_paragraph(line[4:], "Heading 2")
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_runs(paragraph, line[2:])
        else:
            paragraph = document.add_paragraph()
            add_runs(paragraph, line)
        paragraph.paragraph_format.widow_control = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("CampusLoop · 2026-10-09 · ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.author = "CampusLoop team"
    document.core_properties.subject = "八人自选分组、阶段交付与验收任务"
    document.core_properties.title = document.paragraphs[0].text
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)

    # Read the actual saved package to check text, layout settings and structure.
    check = Document(output)
    assert check.paragraphs[0].text == document.paragraphs[0].text
    assert len(check.tables) == 0
    assert len(check.paragraphs) == len(document.paragraphs)
    assert abs(check.sections[0].page_width.cm - 21) < 0.01
    assert abs(check.sections[0].page_height.cm - 29.7) < 0.01
    assert not any("\ufffd" in p.text for p in check.paragraphs)
    print(f"PASS: {output.name}; paragraphs={len(check.paragraphs)}; tables=0; A4")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--wechat-directory", type=Path)
    args = parser.parse_args()
    output_dir = ROOT / "docs" / "management" / "exports"
    for source_name, output_name in SOURCES:
        build(ROOT / "docs" / "management" / source_name, output_dir / output_name)

    if args.wechat_directory:
        target_dir = args.wechat_directory.resolve(strict=True)
        original = target_dir / "任务分配(2).docx"
        if not original.is_file():
            raise FileNotFoundError(original)
        backup = target_dir / "任务分配(2).原始备份-20261009.docx"
        if not backup.exists():
            shutil.copy2(original, backup)
        shutil.copy2(output_dir / SOURCES[0][1], original)
        second = target_dir / "CampusLoop_八人自选分组与个人交付任务书.docx"
        shutil.copy2(output_dir / SOURCES[1][1], second)
        print(f"Backup: {backup}")
        print(f"Updated: {original}")
        print(f"Created: {second}")


if __name__ == "__main__":
    main()
