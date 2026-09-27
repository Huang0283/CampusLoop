from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parents[1] / "CampusLoop_AI_Project_Proposal_English_Template_Adapted.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_table_borders(table, color="808080", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_run_font(run, name="Times New Roman", size=12, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def format_paragraph(paragraph, first_line=True, space_after=7, line=1.08):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.space_after = Pt(space_after)
    if first_line:
        paragraph.paragraph_format.first_line_indent = Inches(0.25)


def add_body(doc, text, first_line=True, space_after=7):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r)
    format_paragraph(p, first_line=first_line, space_after=space_after)
    return p


def add_label_paragraph(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(label)
    set_run_font(r, size=12, bold=True)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 13, bold=True)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(item)
        set_run_font(r)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(f"{index}. {item}")
        set_run_font(r)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, "D9EAF7")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_run_font(r, size=10, bold=True)
    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(value)
            set_run_font(r, size=10)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    set_run_font(run, size=10)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.8)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

normal = doc.styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.08
normal.paragraph_format.space_after = Pt(7)

for style_name in ("List Number", "List Bullet"):
    style = doc.styles[style_name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(12)

add_page_number(section.footer.paragraphs[0])

# Template-style title block.
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(3)
r = p.add_run("Introduction to Software Engineering Assessment")
set_run_font(r, size=18, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(18)
r = p.add_run("CS Group 5 CampusLoop AI Project Proposal")
set_run_font(r, size=17, bold=True)

add_label_paragraph(doc, "Title: ", "CampusLoop AI - An AI-Enhanced Trusted Campus Second-Hand Marketplace")
add_label_paragraph(doc, "Group Number: ", "5")
add_label_paragraph(doc, "Degree Programme: ", "BSc (AI)")
add_label_paragraph(doc, "Team Size: ", "10 members")
add_label_paragraph(doc, "Project Duration: ", "5 weeks (five consecutive one-week delivery phases)")
add_label_paragraph(doc, "Planning Structure: ", "One phase per week; each phase ends with a separately verifiable increment, evidence package, and handoff")
add_label_paragraph(doc, "Date: ", "27 September 2026")

add_heading(doc, "Problem:")
add_body(doc, "University campuses are high-frequency environments for second-hand trading. Textbooks, electronics, furniture, and daily necessities circulate among students, but the transactions are usually scattered across WeChat groups, social feeds, campus forums, and general-purpose marketplaces. These channels are not designed around campus identity, short-distance meetups, or student-to-student trust.")
add_body(doc, "The fragmentation creates inconsistent listing formats, expired information, duplicated conversations, and high search costs. A buyer may describe a need by purpose and budget, such as an affordable device for note-taking, while a seller describes a brand and model. Keyword-only search therefore misses semantically relevant items. At the same time, the lack of reliable condition and price references makes negotiation slow and uncertain.")
add_body(doc, "Campus trading is also time-sensitive. A student may need an item before an examination or a move, while a seller may list it after the buyer has already left campus. When a match is found, the two parties still need to negotiate a place and time, confirm the final arrangement, and coordinate a safe handover. Existing channels provide no consistent scheduling, confirmation, or recovery process for these interactions.")
add_body(doc, "Trust is another major barrier. Misrepresented products, spam, no-shows, duplicate offers, and unresolved disputes are difficult to record and review. Reputation information is often informal and cannot be connected to completed transactions. Finally, AI features in existing products are commonly isolated demonstrations rather than capabilities integrated into listing, matching, pricing, scheduling, and governance workflows.")
add_body(doc, "The project therefore addresses seven connected problems: fragmented information, a semantic gap between needs and listings, pricing asymmetry, temporal mismatch, meetup coordination cost, weak trust and governance, and AI that is disconnected from the transaction workflow.")

add_heading(doc, "Solution:")
add_body(doc, "We plan to develop CampusLoop AI, a web application for a trusted campus second-hand marketplace. The platform combines a standardized transaction lifecycle with explainable AI assistance and human-centered governance. It is designed as a modular proof of concept: the basic marketplace and transaction loop remain usable when AI services are unavailable, while every AI result is presented as assistance rather than an unqualified decision.")
add_body(doc, "The proposed application will implement the following integrated functions:")
add_numbered(doc, [
    "Unified marketplace: standardized product listings, lifecycle states, filtering, favorites, order records, and traceable transaction events.",
    "Semantic search: natural-language queries supported by embedding retrieval with pgvector, combined with keyword search and structured filters. When the model or vector service is unavailable, the system falls back to keyword and filter retrieval.",
    "Wanted marketplace and matching: students can publish persistent requirements. New listings can trigger matching with a relevance score, human-readable reasons, and deduplicated notifications.",
    "Explainable price advice: a rule-based depreciation baseline and a LightGBM regression candidate provide an estimate, a suggested range, influential factors, and data limitations. The result is always labelled as reference information.",
    "Meetup scheduling: the system computes intersections of participants' available time and campus location preferences, recommends feasible options, and supports versioned confirmation and rescheduling.",
    "Trust and risk governance: completed transactions contribute to transparent reputation evidence; rule-based risk signals are routed to human review instead of causing automatic penalties; reports, decisions, and appeals are auditable.",
])
add_body(doc, "The planned architecture uses a React and TypeScript client, a FastAPI service layer, PostgreSQL with pgvector, Redis for optional caching and real-time support, object storage for images and evidence, and background workers for asynchronous tasks. The architecture and interfaces remain modular so that AI components can be replaced or disabled without breaking the basic transaction path.")

add_heading(doc, "Project scope and boundaries")
add_heading(doc, "In-scope capabilities", level=2)
add_numbered(doc, [
    "Identity and access: simulated campus-account registration, login, logout, session handling, profile management, student and administrator roles, field-level visibility, disabled-account handling, and explicit unauthorized outcomes.",
    "Public discovery and authenticated actions: guests may browse public marketplace, product, and wanted-request information; publishing, editing, favouriting, chatting, negotiating, ordering, reviewing, reporting, and viewing private records require an authenticated and authorised user.",
    "Marketplace management: product creation and editing, image metadata, category, condition, price, availability, search, filtering, sorting, pagination, favourites, seller management, and controlled listing lifecycle states.",
    "Wanted marketplace: creation, editing, expiry, closure, browsing, structured requirements, matching results, explanations, and notification control.",
    "Transaction coordination: contextual chat, structured offer and counter-offer, acceptance or rejection, order creation, meetup proposal and versioning, two-party completion confirmation, review eligibility, and a transaction timeline.",
    "Governance: contextual reports for users, products, conversations, and transactions; administrator review; evidence visibility; decision and appeal records; notification read state; and auditable status changes.",
    "Explainable assistance: keyword and semantic retrieval, wanted-to-product matching, price-reference ranges, transparent trust evidence, and risk-review signals, each with a simple baseline, versioned result, explanation, limitation, and non-AI fallback.",
    "Engineering delivery: versioned OpenAPI, generated frontend SDK, database migrations and seed data, repeatable local startup, CI checks, logging, health checks, test evidence, backup and recovery instructions, and release documentation.",
])

add_heading(doc, "Explicitly out of scope", level=2)
add_numbered(doc, [
    "Real-money payment, escrow, refunds, platform commission, delivery, courier integration, and financial guarantees. The project coordinates an offline campus handover and records the agreed transaction state only.",
    "Connection to a university identity provider, government identity verification, face recognition, physical access control, or production student records. Campus verification is simulated for the course demonstration.",
    "Automatic fraud conviction, automatic account punishment, automatic dispute judgment, or AI-controlled order changes. Risk signals support human review and cannot become authoritative facts by themselves.",
    "Unlicensed web scraping, use of real private conversations, or training on personal data without authorisation. Unknown or unavailable datasets are reported as limitations rather than replaced with invented quantities.",
    "A production-scale commercial service-level agreement. Performance and security targets are evaluated in the documented course environment and are not presented as guarantees for unrestricted public deployment.",
])

add_heading(doc, "Users, stakeholders, and principal scenarios")
add_label_paragraph(doc, "Guest visitor: ", "Browses public marketplace and wanted information, opens product details, and is asked to log in only when attempting an identity-dependent action.")
add_label_paragraph(doc, "Student buyer: ", "Searches or describes a need, reviews listing and seller evidence, chats, negotiates, accepts an offer, confirms a meetup version, confirms completion, reviews the seller, and reports a problem when necessary.")
add_label_paragraph(doc, "Student seller: ", "Publishes and manages a listing, responds to questions and offers, accepts one valid offer, coordinates the meetup, confirms completion, reviews the buyer, and manages sold or withdrawn states.")
add_label_paragraph(doc, "Administrator: ", "Views only authorised governance information, reviews reports and risk signals, records a reasoned decision, supports appeal handling, and does not silently alter transaction history.")
add_label_paragraph(doc, "Project operator and QA reviewer: ", "Starts the documented environment, applies migrations, loads synthetic seed data, runs tests and health checks, observes logs, reproduces metrics, verifies backup and recovery, and records the accepted release commit.")
add_body(doc, "The principal end-to-end demonstration begins with a guest browsing public listings. Two students then authenticate as buyer and seller. The seller publishes an item; the buyer finds it through keyword or assisted search; both parties use contextual chat and structured negotiation; acceptance creates an order; they agree and, if necessary, revise a meetup; both confirm completion independently; each becomes eligible to review; notifications reflect the important events; and a contextual report can be routed to an administrator without exposing private evidence to unrelated users.")

add_heading(doc, "Objectives:")
add_heading(doc, "Objective 1: Implement the trusted marketplace and the basic transaction lifecycle", level=2)
add_body(doc, "The first objective is to create a usable campus marketplace and demonstrate the complete non-AI transaction loop. The system must establish user identity and role boundaries before exposing personal, product, order, or report information.")
add_numbered(doc, [
    "Implement a basic login, registration, profile, role, and session model. Students and administrators must have different navigation and permission boundaries, and unauthorized requests must have an explicit result.",
    "Implement product listing, multi-image metadata, category and condition fields, browsing, keyword search, filters, sorting, pagination, favorites, and seller-side listing states such as active, reserved, sold, hidden, and withdrawn.",
    "Implement a wanted marketplace where students can describe a title, budget range, condition requirement, location, and validity period. The flow must support validation and a clear transition from a wanted request to its matching results.",
    "Implement the transaction path from chat and structured offers through counter-offers, acceptance, order creation, meetup arrangement, two-party confirmation, completion, review, and contextual reporting.",
    "Provide a browser interface with responsive navigation, loading, empty, success, failure, and no-permission states. The interface may use static or mock data during the proof-of-concept stage, but it must not present mock data as completed backend integration.",
])
add_body(doc, "Acceptance of Objective 1 will be based on a two-account demonstration: a buyer and a seller browse an item, communicate, negotiate, confirm a meetup version, complete the transaction, and review one another. Order actions must be derived from order state and participant role, and non-participants must not see private conversation or report evidence.")

add_heading(doc, "Objective 2: Integrate explainable AI assistance and campus governance", level=2)
add_body(doc, "The second objective is to connect AI capabilities to real marketplace workflows while preserving transparency, fallback behavior, and human control. Each feature will have a baseline, an evaluation metric, an explanation format, and an explicitly documented unavailable state.")
add_numbered(doc, [
    "Evaluate hybrid semantic search against a keyword baseline using a labelled query set. Report Precision@K, Recall@K, and MRR, and show why a result was retrieved.",
    "Trigger wanted-to-product matching when new listings or wanted requests become eligible. Show the matching factors and relevance score, prevent duplicate notifications, and stop future matching when a request is closed or expired.",
    "Provide price advice with an estimate, interval, factors, sample limitations, and a warning that the seller retains final pricing authority. Compare the candidate model with a median-price baseline using MAE and MAPE where the data supports it.",
    "Recommend meetup options from time intersections and campus locations. Handle no-overlap cases, rescheduling, version increments, and two-party confirmation without silently changing the order fact.",
    "Maintain a transparent trust and risk workflow. Completed trades may contribute verifiable evidence; suspicious activity creates a review case for an administrator; the system must not automatically label a student as fraudulent or impose an irreversible penalty.",
])
add_body(doc, "Acceptance of Objective 2 will include normal operation, unavailable-service fallback, explanation inspection, and privacy checks. AI results must be distinguishable from authoritative transaction facts, and a failed AI service must not prevent listing, browsing, or the basic transaction loop.")

add_heading(doc, "Objective 3: Validate quality, performance, reproducibility, and delivery", level=2)
add_body(doc, "The third objective is to turn the proof of concept into a reproducible software-engineering deliverable. Functional behavior, security boundaries, AI evaluation, and deployment steps must be testable by a non-author of the implementation.")
add_numbered(doc, [
    "Establish unit, API, integration, end-to-end, and permission tests for authentication, listings, search, matching, offers, order state transitions, reviews, reports, and notification behavior.",
    "Test failure and recovery paths including expired offers, duplicate submissions, concurrent acceptance, disconnected chat, message replay, unavailable AI services, invalid files, and unauthorized access.",
    "Run a documented load test with a target of 100 concurrent users and record response-time, error-rate, and resource thresholds. The final threshold and environment must be stated with the raw results.",
    "Provide reproducible setup, migration, seed-data, configuration, CI, logging, backup, and deployment instructions. No secret, real student record, or unlicensed dataset may be required to reproduce the demonstration.",
    "Freeze the scope at the end of Week 4, record implemented, partial, experimental, and deferred features, and deliver a technical report, user guide, administrator guide, and presentation script in Week 5.",
])
add_body(doc, "Acceptance of Objective 3 requires the test commands, data versions, configuration, commit identifier, raw outputs, known limitations, and deployment steps to be available in the repository. A passing result is not claimed for any metric that has not been executed and recorded.")

add_heading(doc, "Benefits:")
add_heading(doc, "5.1 User and community benefits", level=2)
add_body(doc, "Buyers can describe what they need in natural language, find relevant campus items, receive persistent matches, compare transparent price references, and coordinate meetups with less uncertainty. Sellers can reach appropriate buyers, manage listings from one place, understand why a match was suggested, and build a verifiable reputation through completed transactions.")
add_body(doc, "The campus community benefits from a more efficient circular-economy channel. Students can reduce living costs, extend the useful life of goods, and trade in a setting with clearer identity, reporting, and dispute-handling expectations. The platform is designed to encourage responsible peer-to-peer interaction rather than maximize transactions at the expense of safety or autonomy.")
add_heading(doc, "5.2 Technical and engineering benefits", level=2)
add_body(doc, "The project demonstrates separation between a React and TypeScript client and a FastAPI backend, a modular service boundary, state-machine transaction design, optional real-time communication, PostgreSQL and pgvector retrieval, Redis caching, object storage, background jobs, CI, and reproducible testing. It also demonstrates how AI can be integrated into a business workflow with baselines, explanations, fallback paths, and governance instead of being presented as a standalone model demo.")
add_heading(doc, "5.3 Academic and course value", level=2)
add_body(doc, "CampusLoop AI provides practice across the software-engineering lifecycle: requirements, architecture, interface design, implementation, testing, deployment, documentation, and presentation. The ten-member team uses explicit module ownership, five consecutive weekly phase gates, and evidence-based handoffs, producing a traceable case study that can support future campus software projects and demonstrate responsible AI engineering.")

add_heading(doc, "Technical design and integration approach")
add_heading(doc, "Application architecture", level=2)
add_numbered(doc, [
    "Frontend layer: React, TypeScript, Vite, Ant Design, a central route and role configuration, reusable page-state components, generated API types, and a WebSocket or equivalent real-time adapter. The frontend presents state but does not own authoritative permissions, prices, order transitions, or review eligibility.",
    "Service layer: FastAPI modules for identity, profiles, marketplace, wanted requests, conversations, offers, orders, meetup arrangements, reviews, reports, notifications, administration, and intelligent-function orchestration. All state-changing operations validate actor, current state, request version, and idempotency requirements.",
    "Data layer: PostgreSQL stores authoritative users, listings, requests, offers, orders, meetup versions, reviews, reports, notifications, audit fields, and version references. pgvector is optional for semantic retrieval. Database constraints and transactions protect uniqueness, referential integrity, and legal state changes.",
    "Supporting services: object storage holds listing images and authorised evidence; Redis may support caching, deduplication, queues, or real-time coordination; workers execute matching, embedding, evaluation, and notification tasks. The core transaction path has a documented behaviour when an optional service is unavailable.",
    "Intelligent-function layer: M7/M8 services consume versioned business inputs and return versioned suggestions, scores, factors, limitations, and status. M5/M6 control when the service is called, which result is stored, who may see it, and how failure is converted into a safe fallback.",
    "Operations layer: configuration templates, migrations, seed data, CI, structured logs, health checks, data/model version records, backup and restore instructions, and release tags support non-author reproduction.",
])

add_heading(doc, "Core information and state ownership", level=2)
add_body(doc, "The minimum persistent entities are User, Profile, Session, Product, ProductImage, Favourite, WantedRequest, Conversation, Message, Offer, Order, MeetupArrangement, Review, Report, Notification, AiResult, and AuditEvent. Each entity has one service owner, explicit public/private fields, lifecycle states, timestamps, and a traceable relationship to the acting user.")
add_body(doc, "The backend and database are the sources of truth for identity, permissions, listing availability, accepted offer, order state, meetup version, completion confirmations, review eligibility, report status, and administrator decision. Frontend state, cached values, and AI suggestions are derived views. A transition is accepted only when the actor, previous state, version, and domain invariant permit it; otherwise the API returns a stable error that the client can explain.")

add_heading(doc, "API, SDK, and change control", level=2)
add_body(doc, "The OpenAPI document is the shared contract for HTTP endpoints, schemas, enumerations, error responses, authentication requirements, and examples. The frontend SDK is generated from the committed contract and is not manually treated as a separate source of truth. Contract changes require a versioned pull request, impact notes, regenerated SDK output, affected tests, and confirmation from the provider, consumer, M10, and M1 when scope changes.")

add_heading(doc, "Non-functional requirements")
add_heading(doc, "Security and privacy", level=2)
add_numbered(doc, [
    "Passwords are never stored or logged in plain text; tokens and session identifiers are protected, expire, and can be invalidated. Secrets are supplied through ignored environment files or secure CI variables.",
    "Every private endpoint enforces server-side role, ownership, participant, and state checks. Public responses exclude email, token, private timestamps, conversation content, report evidence, and internal governance fields unless the requester is authorised.",
    "Uploads are validated by type, size, count, ownership, and storage key. Logs avoid credentials and sensitive content. Synthetic accounts and data are used for demonstrations and automated tests.",
    "Reports, administrator actions, and AI risk signals are auditable. A user-facing result distinguishes a report allegation, an automated signal, and a human decision.",
])

add_heading(doc, "Reliability and consistency", level=2)
add_numbered(doc, [
    "Critical writes use transactions, uniqueness rules, idempotency keys or equivalent safeguards, and optimistic or pessimistic concurrency control where duplicate acceptance or confirmation is possible.",
    "Reconnect and retry behaviour must not create duplicate messages, offers, orders, reviews, reports, or notifications. The client can reload the current authoritative state after interruption.",
    "Optional AI, cache, queue, vector, or object-storage failures have explicit user-visible outcomes and recovery procedures. AI failure does not block public browsing or the basic transaction loop.",
])

add_heading(doc, "Usability and accessibility", level=2)
add_numbered(doc, [
    "The interface supports desktop and practical mobile widths, stable navigation, keyboard access to primary actions, readable contrast, meaningful labels, confirmation for destructive actions, and consistent loading, empty, success, failure, retry, and no-permission states.",
    "Role and action visibility is predictable: hidden navigation is not used as the only security mechanism, unavailable actions explain why, and login prompts appear at the attempted protected action rather than blocking permitted public browsing.",
    "AI results identify the assisting feature, important factors, confidence or limitation, fallback state, and the fact that the user or authorised reviewer retains decision authority.",
])

add_heading(doc, "Performance and reproducibility", level=2)
add_numbered(doc, [
    "The final release includes a documented 100-concurrent-user load scenario. The environment, endpoint mix, data volume, duration, response-time distribution, error rate, resource use, and raw output are recorded before any performance claim is made.",
    "A non-author can install dependencies, configure services, apply migrations, load synthetic seed data, start the system, run core tests, reproduce accepted AI metrics, and follow the main demonstration from committed instructions.",
    "Build, lint, unit, API, contract, integration, E2E, permission, recovery, and evaluation commands record the exact commit and relevant data or model version.",
])

add_heading(doc, "Validation and acceptance strategy")
add_body(doc, "Acceptance follows a traceable evidence chain: requirement ID -> design or contract -> implementation commit -> automated or manual test -> recorded result -> phase handoff -> final release claim. The author of a feature performs initial checks, another group member reviews the change, and M10 conducts independent phase acceptance on the integrated commit. M1 confirms that the demonstrated result remains within the approved scope.")
add_numbered(doc, [
    "Functional acceptance: normal, validation, permission, state, duplicate, concurrency, interruption, and recovery scenarios for every MVP journey.",
    "Contract acceptance: OpenAPI validation, generated SDK consistency, migration and seed execution, service health, success/error examples, and page-to-API mapping.",
    "Intelligent-function acceptance: documented dataset or test set, simple baseline, candidate command, raw metrics, failure analysis, explanation inspection, privacy check, version record, and unavailable-service fallback.",
    "Release acceptance: clean installation, final regression, P0/P1 defect status, backup and restore, load-test evidence, documentation reconciliation, demonstration rehearsal, final tag, and submission checklist.",
])

add_heading(doc, "Key risks and mitigation")
add_numbered(doc, [
    "Scope overload within five weeks: protect the non-AI transaction loop, remove Stretch before Core, freeze contracts in Week 2 and features in Week 4, and require M1 approval for every scope change.",
    "Cross-group dependency or merge conflict: assign one owner to shared contracts and files, publish providers before consumers integrate, use small task pull requests, prohibit force-push on shared branches, and merge each group through the phase integration branch.",
    "Insufficient or unlicensed AI data: use rule and keyword baselines, record unknowns honestly, restrict claims to executed evidence, and mark unsupported capabilities conditional or no-go.",
    "Permission or privacy leakage: maintain field-visibility and role matrices, enforce checks server-side, use synthetic data, test non-participant access, and independently review report and conversation evidence exposure.",
    "Inconsistent transaction state under duplicate or concurrent actions: define legal state machines and invariants, protect writes with transaction/idempotency controls, and test simultaneous acceptance, confirmation, review, and retry scenarios.",
    "Environment or optional-service failure: version dependencies and configuration, provide health checks and deterministic fallbacks, retain logs, verify clean startup, and test backup, restore, reconnect, and service degradation.",
    "Unreproducible performance or AI claims: preserve commands, parameters, random seeds, data/model versions, raw outputs, hardware/environment notes, and the exact release commit.",
])

add_heading(doc, "Expected final deliverables")
add_numbered(doc, [
    "A tagged source repository containing the frontend, backend, intelligent-function code, configuration templates, migrations, seed data, generated SDK, automated tests, and CI definitions.",
    "A deployable demonstration environment with startup, health, shutdown, migration, backup, restore, and troubleshooting instructions that do not require a secret or private dataset from an individual member.",
    "Requirements, architecture, database, OpenAPI, state-machine, security/privacy, AI design, test strategy, risk, handoff, and known-limitation documentation connected to the final code version.",
    "Raw and summarised functional, permission, integration, E2E, recovery, performance, and AI evaluation evidence, including failures and deferred items.",
    "A user guide, administrator guide, technical report, contribution record, presentation slides or script, reproducible demonstration accounts, final checklist, and accepted release tag.",
])

add_heading(doc, "Timeline")
add_body(doc, "The project is organised as five consecutive one-week phases. A phase is not merely a time period: it must produce a reviewable increment, repository evidence, an acceptance result, and an explicit handoff to the following week. Work that misses a Friday gate is recorded with an owner, impact, recovery action, and revised date; it is not silently treated as complete or allowed to invalidate the next phase.")

add_heading(doc, "Phase 1 / Week 1: Scope, feasibility, and planning baseline", level=2)
add_label_paragraph(doc, "Weekly outcome: ", "The team agrees what the project will and will not deliver, who owns each capability, what the principal user journeys are, and whether the proposed platform and AI functions are feasible enough to continue.")
add_label_paragraph(doc, "Required work: ", "M1 and M10 create the scope, priority, traceability, risk, test, and defect-management baseline. M2-M4 audit the frontend and produce the page inventory, role navigation, low-fidelity user journeys, and page-state rules. M5, M6, and M9 define account boundaries, domain entities, state-machine candidates, system context, ER candidates, and runtime constraints. M7 and M8 define keyword/search, matching, price, trust, and risk baselines with data-source and licensing checks.")
add_label_paragraph(doc, "Mandatory evidence: ", "Requirement IDs and priorities; RACI and escalation path; page and route inventory; transaction swimlane; domain/entity catalogue; state and permission candidates; system context and ER diagram; AI capability feasibility files; test strategy; review log; first-presentation rehearsal record.")
add_label_paragraph(doc, "Exit gate: ", "Every in-scope capability has one owner, a planned phase, an acceptance route, and no unresolved responsibility gap. The four Phase 1 group issues are reviewed, their summaries are integrated, and M1/M10 record accepted limitations and blockers.")
add_label_paragraph(doc, "Handoff to Week 2: ", "Accepted candidate requirements, page journeys, domain rules, data feasibility conclusions, test scenarios, and open decisions become versioned inputs to the design and contract baseline.")

add_heading(doc, "Phase 2 / Week 2: Design, contracts, and runnable foundation", level=2)
add_label_paragraph(doc, "Weekly outcome: ", "The team freezes a mutually consistent design baseline that a non-author can start, inspect, and use for implementation without inventing fields, permissions, states, or AI outputs.")
add_label_paragraph(doc, "Required work: ", "M2-M4 convert approved journeys into clickable prototypes, reusable component specifications, route-role rules, and page-to-API mappings. M5/M6 publish the versioned OpenAPI contract, success and failure examples, permission rules, database schema, migrations, seed data, and minimal service skeleton. M9 supplies the reproducible local environment, port and environment-variable contract, CI smoke checks, and storage/service health checks. M7/M8 publish dataset, labelling, evaluation, output-schema, explanation, and fallback contracts. M10 links contract tests and acceptance scenarios to requirement IDs.")
add_label_paragraph(doc, "Mandatory evidence: ", "Clickable prototype; design/component specification; OpenAPI file and generated SDK; database migration and seed result; environment start command and health output; AI sample/evaluation contract; page/API mapping; contract-review record; Phase 2 verification and handoff files.")
add_label_paragraph(doc, "Exit gate: ", "The unified environment starts from documented commands; the API specification validates; migrations apply; the generated SDK matches the committed contract; prototypes demonstrate both success and error states; M1 freezes scope and M10 confirms that the contracts are independently testable.")
add_label_paragraph(doc, "Handoff to Week 3: ", "Implementation owners receive frozen interface versions, database versions, test data, page mappings, AI service schemas, and executable acceptance scenarios.")

add_heading(doc, "Phase 3 / Week 3: Real MVP implementation and end-to-end transaction loop", level=2)
add_label_paragraph(doc, "Weekly outcome: ", "Two test accounts complete the real persisted marketplace and transaction journey through the browser using actual APIs rather than fixed responses or unlabelled mocks.")
add_label_paragraph(doc, "Required work: ", "M5 implements authentication, sessions, profile and role boundaries. M6 implements products, wanted requests, search, favourites, chat/message transport, structured offers, orders, meetup versions, two-party completion, reviews, reports, and notifications with validated state transitions and idempotency. M9 supports migrations, seed accounts, observability, storage, and repeatable startup. M7/M8 expose the minimum callable keyword/matching and rule-based price/trust baselines required by the frozen contracts. M2-M4 replace Phase 2 mocks with real SDK calls and implement loading, empty, error, retry, permission, and reconnect behaviour. M10 executes API, permission, state-machine, and browser E2E tests.")
add_label_paragraph(doc, "Mandatory evidence: ", "Running services; committed migrations and seed data; API and WebSocket test output; frontend build/lint results; two-account E2E record covering registration/login, listing, browsing/search, chat, offer, order, meetup, both confirmations, review, notification, and report; defect list and tagged MVP commit.")
add_label_paragraph(doc, "Exit gate: ", "The transaction loop survives refresh and restart, private data is inaccessible to non-participants, invalid transitions and duplicate submissions are rejected consistently, and no required E2E step depends on an unlabelled mock.")
add_label_paragraph(doc, "Handoff to Week 4: ", "The accepted MVP version, real domain events, current defects, production-shaped data structures, and performance baseline become the integration target for intelligent functions and governance.")

add_heading(doc, "Phase 4 / Week 4: AI integration, governance, and feature freeze", level=2)
add_label_paragraph(doc, "Weekly outcome: ", "At least two intelligent capabilities are integrated into the real workflow, measured against simple baselines, explained to users, and proven not to block the marketplace when unavailable. The feature set is frozen at the end of the week.")
add_label_paragraph(doc, "Required work: ", "M7 integrates and evaluates hybrid search and wanted-to-product matching. M8 integrates the price-advice candidate and the agreed trust/risk assistance that has sufficient evidence. M5/M6 enforce permissions, persist versioned AI results, orchestrate background triggers, deduplicate notifications, and preserve authoritative transaction facts. M2-M4 display explanations, limitations, unavailable states, manual alternatives, and the minimum administrator review workflow. M9 records model/data versions and background-service health. M10 runs baseline comparison, privacy, permission, degradation, and recovery tests; M1 approves the freeze list.")
add_label_paragraph(doc, "Mandatory evidence: ", "Versioned dataset or test set; baseline and candidate commands; raw metric outputs; explanation examples; API/UI integration evidence; unavailable-service demonstration; privacy and governance checks; freeze list dividing implemented, experimental, deferred, and removed functions.")
add_label_paragraph(doc, "Exit gate: ", "Reported metrics are reproducible from committed instructions, explanations match actual inputs, AI outputs cannot directly alter orders or penalties, fallback paths preserve the basic transaction loop, critical defects have owners, and no unapproved feature remains open after Friday.")
add_label_paragraph(doc, "Handoff to Week 5: ", "The frozen release candidate, metric report, model/data versions, defect priorities, known limitations, and demonstration scope become the sole basis for final hardening and submission.")

add_heading(doc, "Phase 5 / Week 5: Stabilisation, reproducibility, and final delivery", level=2)
add_label_paragraph(doc, "Weekly outcome: ", "A non-author can deploy the frozen project in a clean environment, run the documented tests, recover from defined failures, and demonstrate the agreed user and administrator journeys using the final tagged version.")
add_label_paragraph(doc, "Required work: ", "All implementation owners fix approved defects without adding scope. M9 verifies clean deployment, configuration, migration, backup, restore, logs, health checks, and release packaging. M10 executes full regression, permission/security checks, failure recovery, and the documented 100-concurrent-user load test. M1 reconciles requirements with implemented evidence and compiles the technical report, contribution record, user guide, administrator guide, presentation, and final submission checklist. Every owner verifies that documentation, screenshots, metrics, and demonstrations reference the same release commit.")
add_label_paragraph(doc, "Mandatory evidence: ", "Clean-install record; final CI and regression output; security/permission checks; load-test configuration and raw result; backup/restore evidence; known-issues register; user and administrator guides; technical report; presentation script; contribution record; final tag and signed submission checklist.")
add_label_paragraph(doc, "Exit gate: ", "No open release-blocking defect remains, the clean deployment and main demonstration pass, repository and written claims agree, secrets and personal data are absent, all required files are present, and M1/M10 sign the final acceptance record.")
add_label_paragraph(doc, "Final result: ", "The final tag, source, reproducibility package, test evidence, documentation, and presentation materials form one auditable submission package.")

doc.add_page_break()
add_heading(doc, "Five-phase delivery summary", level=2)
add_table(doc, ["Phase / Week", "Verified increment", "Mandatory evidence", "Exit milestone"], [
    ("P1 / W1", "Agreed scope, ownership, journeys, domain boundaries, AI feasibility, and quality plan", "Requirements/traceability, wireframes, state and data candidates, feasibility and review records", "M1: Scope and feasibility baseline accepted"),
    ("P2 / W2", "Clickable design and frozen API/database/environment/AI/test contracts", "Prototype, OpenAPI + SDK, migration/seed, startup proof, contract and handoff records", "M2: Design and runnable foundation accepted"),
    ("P3 / W3", "Persisted two-account MVP transaction loop using real services", "API/E2E output, state/permission tests, defect list, tagged MVP commit", "M3: Real MVP loop accepted"),
    ("P4 / W4", "Measured AI integration, explanations, fallback, governance, and feature freeze", "Metric outputs, integration/degradation evidence, privacy checks, freeze list", "M4: Integrated release candidate frozen"),
    ("P5 / W5", "Clean deployment, regression, recovery, documentation, and submission package", "CI/regression/load/recovery evidence, guides, report, presentation, final tag", "M5: Final delivery accepted"),
], widths=[0.9, 2.4, 2.8, 1.6])

add_heading(doc, "Action plan")
add_body(doc, "The ten-person team is organised into four delivery groups: management and quality (M1/M10), frontend experience (M2-M4), backend and platform (M5/M6/M9), and intelligent functions (M7/M8). Each person owns a bounded deliverable, while cross-group work follows a provider-to-consumer-to-independent-acceptance sequence. A page cannot prove that an API exists, a mock cannot close a real-integration task, and an AI output cannot be accepted without its data, baseline, metric, explanation, and fallback evidence.")

add_heading(doc, "Member responsibilities across the five weeks", level=2)
add_table(doc, ["Member", "Primary ownership", "Weekly delivery responsibility"], [
    ("M1", "Scope, requirements, planning, integration governance", "W1 scope/RACI/traceability; W2 baseline freeze; W3 MVP scope acceptance; W4 feature freeze; W5 final reconciliation, report, presentation, and submission sign-off."),
    ("M2", "Frontend lead, shell, auth/profile, shared components", "W1 audit/navigation/design rules; W2 clickable shell and API map; W3 real auth/profile integration; W4 administrator/AI states and freeze; W5 accessibility, regression, and frontend delivery notes."),
    ("M3", "Marketplace, wanted, search/matching/price UI", "W1 journeys and field inventory; W2 clickable prototypes and contracts; W3 real marketplace/wanted integration; W4 search/matching/price explanations and fallback; W5 regression and user-guide evidence."),
    ("M4", "Chat, offers, orders, meetup, reviews, reports, notifications", "W1 transaction wireframes and state actions; W2 clickable flow and API mapping; W3 real two-account transaction integration; W4 governance/error-state freeze; W5 E2E regression and demonstration script."),
    ("M5", "Authentication, sessions, profile, roles, governance permissions", "W1 identity/visibility/risk rules; W2 auth contract and service skeleton; W3 real auth/profile/permission APIs; W4 administration and AI privacy controls; W5 security, recovery, and deployment verification."),
    ("M6", "Marketplace and transaction domain/API/WebSocket", "W1 entities, states, invariants, concurrency risks; W2 OpenAPI/database contract; W3 real marketplace/chat/order implementation; W4 AI orchestration and persisted results; W5 state, concurrency, and recovery regression."),
    ("M7", "Search, semantic retrieval, matching, evaluation", "W1 feasibility and rule baseline; W2 dataset/evaluation/output contract; W3 callable keyword/matching baseline; W4 hybrid integration and measured comparison; W5 reproducibility package and limitation statement."),
    ("M8", "Price advice, trust, risk assistance, evaluation", "W1 data/permission feasibility and rule baseline; W2 dataset/output/evaluation contract; W3 callable rule services; W4 supported model/rule integration and governance evaluation; W5 reproducibility and limitation statement."),
    ("M9", "Database, environments, storage, jobs, CI/CD, operations", "W1 system/ER/runtime plan; W2 migrations, seed, ports, environment and CI smoke checks; W3 stable integration environment and observability; W4 model/job versioning; W5 clean deployment, backup/restore, logs, and release package."),
    ("M10", "Independent QA, defects, evidence, release acceptance", "W1 test/defect baseline; W2 contract and environment tests; W3 API/permission/state/E2E acceptance; W4 metric, privacy, fallback and freeze verification; W5 full regression, load/recovery tests, and final sign-off."),
], widths=[0.6, 2.2, 4.9])

add_heading(doc, "Weekly operating and merge discipline", level=2)
add_numbered(doc, [
    "Monday - input lock and planning: the phase integration branch is created from the previously accepted main commit; every task records its owner, input version, output path, reviewer, validation command, and Friday acceptance result.",
    "Tuesday to Wednesday - bounded implementation: providers publish reviewed contracts or services before consumers integrate them. Shared files have one owner and an agreed merge order; direct commits to main and force-pushes to shared branches are prohibited.",
    "Thursday - group integration: personal task pull requests are reviewed into the relevant group branch, group owners create summary pull requests to the phase integration branch, and M10 begins independent verification on the integrated commit.",
    "Friday - acceptance and handoff: M10 records executed checks and defects; M1 confirms scope and unresolved items; one phase-closure pull request merges the accepted integration commit to main and records the next phase's exact inputs.",
    "Incomplete work: any missed item is recorded as blocker / owner / dependency / impact / next action / exact date. It is either formally removed from the phase or remains open; it cannot be marked complete through screenshots, verbal confirmation, mock data, or an unmerged branch.",
])

add_heading(doc, "Five-week action register", level=2)
add_table(doc, ["Deadline", "Required coordinated action", "Acceptance responsibility", "Completion evidence"], [
    ("End of W1", "Complete and cross-review the four scope/feasibility work packages; reconcile role, page, domain, data, risk, and test candidates.", "M1 scope; M10 testability; group leads content", "Four group summaries, review/rehearsal records, phase handoff, accepted integration commit"),
    ("End of W2", "Freeze requirements and versioned UI/API/database/environment/AI/test contracts; demonstrate a clean startup and clickable prototype.", "M1 baseline; M10 contract/startup tests", "Validated OpenAPI and SDK, migration/seed proof, prototype, contract review, handoff"),
    ("End of W3", "Implement and integrate the real persisted two-account marketplace and transaction loop; close or assign all critical defects.", "M10 E2E/permission/state tests; M1 MVP scope", "Running system, test output, defect log, MVP tag, two-account acceptance record"),
    ("End of W4", "Integrate at least two justified intelligent functions; run baseline comparison, explanation, privacy and fallback checks; freeze scope.", "M10 metrics/degradation; M1 freeze approval", "Raw metrics, integration proof, governance checks, freeze list, release-candidate commit"),
    ("End of W5", "Perform clean deployment, regression, security, load, backup/restore, documentation reconciliation, demonstration, and packaging.", "M10 release tests; M1 final sign-off", "Final tag, release evidence, guides/report/presentation, contribution record, signed checklist"),
], widths=[1.0, 3.3, 1.6, 2.0])

add_heading(doc, "Note")
add_body(doc, "This proposal is a preliminary research and planning document. Features, technical approaches, evaluation datasets, and dates remain subject to refinement after Week 1 user research and supervisor feedback. The project will follow academic-integrity requirements: external sources will be cited, personal data will not be used without authorization, and AI tools will be disclosed and used according to university guidance. Detailed requirements, system design, database schema, AI specifications, testing strategies, issue records, and risk registers are maintained in the accompanying CampusLoop AI repository documentation.")

doc.core_properties.title = "CampusLoop AI Project Proposal"
doc.core_properties.subject = "Introduction to Software Engineering Assessment"
doc.core_properties.author = "CS Group 5"
doc.core_properties.keywords = "CampusLoop AI, software engineering, project proposal"
doc.save(OUT)
print(OUT)
